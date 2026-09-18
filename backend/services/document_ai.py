"""Document ingestion and text extraction for common GS420 formats."""
from pathlib import Path
import csv
class DocumentAI:
    TEXT_EXT={".txt",".md",".csv",".pdf",".docx",".xlsx"}
    def extract(self,path):
        p=Path(path);ext=p.suffix.lower()
        if ext==".pdf": return self._pdf(p)
        if ext==".docx": return self._docx(p)
        if ext==".xlsx": return self._xlsx(p)
        if ext==".csv": return self._csv(p)
        return {"text":p.read_text(encoding="utf-8",errors="ignore"),"pages":1,"type":ext}
    def _pdf(self,p):
        from pypdf import PdfReader
        pages=PdfReader(str(p)).pages
        return {"text":"\n".join(x.extract_text() or "" for x in pages),"pages":len(pages),"type":".pdf"}
    def _docx(self,p):
        from docx import Document
        d=Document(str(p));return {"text":"\n".join(x.text for x in d.paragraphs),"pages":1,"type":".docx"}
    def _xlsx(self,p):
        from openpyxl import load_workbook
        w=load_workbook(p,read_only=True,data_only=True)
        rows=[]
        for s in w.worksheets:
            for row in s.iter_rows(values_only=True):rows.append("\t".join("" if v is None else str(v) for v in row))
        return {"text":"\n".join(rows),"pages":1,"type":".xlsx"}
    def _csv(self,p):
        with p.open(encoding="utf-8-sig",newline="",errors="ignore") as f:return {"text":"\n".join("\t".join(r) for r in csv.reader(f)),"pages":1,"type":".csv"}
