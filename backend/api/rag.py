"""RAG document extraction and search API."""
from __future__ import annotations
import io, os, tempfile
from typing import Any
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from backend.rag import DocumentStore, EmbeddingProvider
from backend.vision import OCRService
from backend.config import get_settings
from backend.voter_ai import VoterStore, extract_and_store

router=APIRouter(prefix="/api/rag",tags=["rag"])
settings=get_settings()
store=DocumentStore(embedder=EmbeddingProvider(settings.embedding_model_id) if settings.embedding_model_id else None)
ocr=OCRService()
voters=VoterStore()

class SearchRequest(BaseModel):
 query:str=Field(min_length=1,max_length=5000)
 limit:int=Field(default=5,ge=1,le=20)
 semantic:bool=False

def _extract_text(filename,data,ocr_scanned=False,language="ben+eng",preprocess="balanced"):
 suffix=os.path.splitext(filename)[1].lower()
 if suffix in {".txt",".md",".csv",".json"}: return data.decode("utf-8",errors="replace")
 if suffix==".docx":
  from docx import Document
  return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
 if suffix==".pdf":
  from pypdf import PdfReader
  text="\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages).strip()
  if text or not ocr_scanned:return text
  import pypdfium2 as pdfium
  with tempfile.TemporaryDirectory() as td:
   path=os.path.join(td,"input.pdf");open(path,"wb").write(data);pdf=pdfium.PdfDocument(path);pages=[]
   for i,page in enumerate(pdf):
    image=page.render(scale=2).to_pil(); ip=os.path.join(td,f"{i}.png");image.save(ip)
    pages.append(ocr.extract_text(ip,language,preprocess)["text"])
   return "\n".join(pages)
 raise ValueError("Supported formats: TXT, MD, CSV, JSON, PDF, DOCX.")

@router.get("/stats")
def stats(): return store.stats()

@router.post("/documents")
async def upload_document(file:UploadFile=File(...),embed:bool=Form(False),ocr_scanned:bool=Form(False),language:str=Form("ben+eng"),preprocess:str=Form("balanced"),extract_voters:bool=Form(False)):
 data=await file.read()
 if not data:raise HTTPException(400,"Document is empty.")
 if len(data)>20*1024*1024:raise HTTPException(413,"Document is too large.")
 try:
  filename=file.filename or "document";text=_extract_text(filename,data,ocr_scanned,language,preprocess)
  if not text.strip():raise ValueError("No extractable text found.")
  result=store.add(filename,text,create_embeddings=embed)
  if extract_voters:
   result["voters"]=extract_and_store(result["id"],text,voters)
  return result
 except RuntimeError as e:raise HTTPException(503,str(e))
 except ValueError as e:raise HTTPException(415,str(e))

@router.post("/search")
def search(req:SearchRequest):
 if req.semantic and store.embedder is None:raise HTTPException(503,"Semantic search is not configured.")
 return {"results":store.search(req.query,req.limit,req.semantic)}

@router.get("/documents/{document_id}")
def get_document(document_id:str):
 item=store.get(document_id)
 if not item:raise HTTPException(404,"Document not found.")
 return item

@router.delete("/documents/{document_id}")
def delete_document(document_id:str):return {"deleted":store.delete(document_id)}
